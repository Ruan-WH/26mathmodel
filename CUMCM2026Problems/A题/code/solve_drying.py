"""Numerical solution for CUMCM 2026 Problem A.

The model uses a node-centred finite-volume discretisation in the radial
coordinate, fully implicit diffusion, Picard iteration for nonlinear
properties, and a material-coordinate transform for uniform radial shrinkage.
"""

from __future__ import annotations

import argparse
import json
import math
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator
from scipy.linalg import solve_banded
from scipy.special import expi


ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = ROOT / "附件"
RESULTS_DIR = ROOT / "results"
H_T = 25.0
H_M = 8.0e-7
R_FIXED = 0.02
NODES = 321
GRID_POWER = 2.0
INITIAL_T = 28.0
INITIAL_C = 2.55
THRESHOLD_C = 0.15


@dataclass
class Environment:
    time_s: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray
    temp_interp: PchipInterpolator
    moisture_interp: PchipInterpolator

    def values(
        self,
        time_s: float,
        terminal_temp_scale: float = 1.0,
        terminal_moisture_scale: float = 1.0,
    ) -> tuple[float, float]:
        cutoff = float(self.time_s[-1])
        if time_s <= cutoff:
            return float(self.temp_interp(time_s)), float(self.moisture_interp(time_s))
        return (
            float(self.temperature_c[-1] * terminal_temp_scale),
            float(self.moisture[-1] * terminal_moisture_scale),
        )


@dataclass
class RadiusHistory:
    time_s: np.ndarray
    radius_m: np.ndarray
    interp: PchipInterpolator
    derivative: PchipInterpolator
    monotone_adjustments: int

    def values(self, time_s: float) -> tuple[float, float]:
        bounded = min(max(time_s, float(self.time_s[0])), float(self.time_s[-1]))
        return float(self.interp(bounded)), float(self.derivative(bounded))


def _read_numeric_sheet(path: Path) -> np.ndarray:
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book.active
    rows: list[list[float]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        rows.append([float(value) for value in row])
    return np.asarray(rows, dtype=float)


def load_environment() -> Environment:
    data = _read_numeric_sheet(INPUT_DIR / "附件1.xlsx")
    time_s, temperature_c, moisture = data.T
    if not np.all(np.diff(time_s) > 0):
        raise ValueError("附件1时间列不是严格递增。")
    if not np.all(np.isfinite(data)):
        raise ValueError("附件1含非有限数。")
    return Environment(
        time_s=time_s,
        temperature_c=temperature_c,
        moisture=moisture,
        temp_interp=PchipInterpolator(time_s, temperature_c),
        moisture_interp=PchipInterpolator(time_s, moisture),
    )


def load_radius_history() -> RadiusHistory:
    data = _read_numeric_sheet(INPUT_DIR / "附件2.xlsx")
    time_s = data[:, 0]
    raw_radius_m = data[:, 1] / 100.0
    if not np.all(np.diff(time_s) > 0):
        raise ValueError("附件2时间列不是严格递增。")
    monotone_radius_m = np.minimum.accumulate(raw_radius_m)
    adjustments = int(np.count_nonzero(np.abs(monotone_radius_m - raw_radius_m) > 1e-12))
    interp = PchipInterpolator(time_s, monotone_radius_m, extrapolate=False)
    return RadiusHistory(
        time_s=time_s,
        radius_m=monotone_radius_m,
        interp=interp,
        derivative=interp.derivative(),
        monotone_adjustments=adjustments,
    )


def property_q1(temperature_c: np.ndarray, moisture: np.ndarray) -> tuple[np.ndarray, ...]:
    del temperature_c
    rho = np.full_like(moisture, 820.0)
    cp = np.full_like(moisture, 2600.0)
    k = np.full_like(moisture, 0.36)
    safe_c = np.maximum(moisture, 1e-8)
    diffusivity = 7.0e-9 * np.exp(-0.89 / safe_c)
    return rho, cp, k, diffusivity


def property_q23(temperature_c: np.ndarray, moisture: np.ndarray) -> tuple[np.ndarray, ...]:
    safe_c = np.maximum(moisture, 1e-8)
    rho = 650.0 + 128.0 * safe_c
    cp = 1450.0 + 2736.0 * safe_c / (safe_c + 1.0)
    k = 0.21 + 0.38 * safe_c / (safe_c + 1.0)
    diffusivity = (
        2.4e-3
        * np.exp(-0.45 / safe_c)
        * np.exp(-3850.0 / (temperature_c + 273.15))
    )
    return rho, cp, k, diffusivity


def property_q4(temperature_c: np.ndarray, moisture: np.ndarray) -> tuple[np.ndarray, ...]:
    safe_c = np.maximum(moisture, 1e-8)
    rho = 760.0 + 90.0 * safe_c
    cp = 1850.0 + 2150.0 * safe_c / (safe_c + 1.0)
    k = 0.12 + 0.20 * safe_c / (safe_c + 1.0)
    diffusivity = (
        4.2e-4
        * np.exp(-0.30 / safe_c)
        * np.exp(-3850.0 / (temperature_c + 273.15))
    )
    return rho, cp, k, diffusivity


def face_interpolate(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    """Linear face interpolation for a smooth state-dependent coefficient."""
    return 0.5 * (left + right)


def moisture_faces(temperature: np.ndarray, moisture: np.ndarray,
                   property_model: Callable, axis: int = 0) -> np.ndarray:
    """Kirchhoff secant flux: integrate D(T_face,C) over the face concentration jump.

    F_a(C) = C exp(-a/C) + a Ei(-a/C), with F_a' = exp(-a/C).
    This avoids arithmetic averaging across the strongly nonlinear dry surface layer.
    """
    left = np.take(moisture, np.arange(moisture.shape[axis] - 1), axis=axis)
    right = np.take(moisture, np.arange(1, moisture.shape[axis]), axis=axis)
    tl = np.take(temperature, np.arange(temperature.shape[axis] - 1), axis=axis)
    tr = np.take(temperature, np.arange(1, temperature.shape[axis]), axis=axis)
    midpoint = np.maximum(0.5 * (left + right), 1e-8)
    exponent = getattr(property_model, 'moisture_exponent', 0.45)
    def primitive(c):
        c = np.maximum(c, 1e-8)
        return c * np.exp(-exponent / c) + exponent * expi(-exponent / c)
    jump = right - left
    close = np.abs(jump) < 1e-5 * midpoint
    secant = np.exp(-exponent / midpoint)
    np.divide(primitive(right) - primitive(left), jump,
              out=secant, where=~close)
    prefactor = property_model(0.5 * (tl + tr), np.ones_like(midpoint))[3] * np.exp(exponent)
    return prefactor * np.maximum(secant, 0.0)


property_q1.moisture_exponent = 0.89
property_q23.moisture_exponent = 0.45
property_q4.moisture_exponent = 0.30


def node_control_volumes(coordinate: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    """Return dimensionless/physical radial CV volumes and internal face radii."""
    faces = 0.5 * (coordinate[:-1] + coordinate[1:])
    lower = np.concatenate(([0.0], faces))
    upper = np.concatenate((faces, [coordinate[-1]]))
    volumes = math.pi * (upper**2 - lower**2)
    return volumes, faces


def radial_grid(nodes: int) -> np.ndarray:
    """Surface-graded material grid; resolve the initial sub-millimetre dry layer."""
    return 1.0 - (1.0 - np.linspace(0.0, 1.0, nodes)) ** GRID_POWER


def implicit_radial_step(
    old: np.ndarray,
    storage: np.ndarray,
    conductivity: np.ndarray,
    dt: float,
    coordinate: np.ndarray,
    diffusion_scale: float,
    boundary_conductance: float,
    boundary_value: float,
    advection_velocity: np.ndarray | None = None,
    face_conductivity: np.ndarray | None = None,
) -> np.ndarray:
    """One implicit node-centred radial FVM step on coordinate [0, 1] or [0, R]."""
    count = len(old)
    volumes, faces = node_control_volumes(coordinate)
    delta = np.diff(coordinate)
    face_property = (face_interpolate(conductivity[:-1], conductivity[1:])
                     if face_conductivity is None else face_conductivity)
    conductance = 2.0 * math.pi * faces * face_property * diffusion_scale / delta

    capacity = storage * volumes / dt
    diagonal = capacity.copy()
    lower = np.zeros(count - 1)
    upper = np.zeros(count - 1)
    rhs = capacity * old

    diagonal[:-1] += conductance
    diagonal[1:] += conductance
    upper -= conductance
    lower -= conductance

    diagonal[-1] += boundary_conductance
    rhs[-1] += boundary_conductance * boundary_value

    if advection_velocity is not None:
        # PDE form: phi_t + u(xi) phi_xi = diffusion; u >= 0 under shrinkage.
        advective = storage[1:] * volumes[1:] * np.maximum(advection_velocity[1:], 0.0) / delta
        diagonal[1:] += advective
        lower -= advective

    matrix = np.zeros((3, count))
    matrix[0, 1:] = upper
    matrix[1, :] = diagonal
    matrix[2, :-1] = lower
    return solve_banded((1, 1), matrix, rhs, check_finite=False)


def advance_coupled(
    temperature_old: np.ndarray,
    moisture_old: np.ndarray,
    dt: float,
    radius_m: float,
    radius_rate_m_s: float,
    ambient_temperature_c: float,
    ambient_moisture: float,
    property_model: Callable[[np.ndarray, np.ndarray], tuple[np.ndarray, ...]],
    moving_domain: bool,
    tolerance_temperature: float = 1e-8,
    tolerance_moisture: float = 1e-10,
    max_iterations: int = 30,
) -> tuple[np.ndarray, np.ndarray, int]:
    temperature = temperature_old.copy()
    moisture = moisture_old.copy()
    if moving_domain:
        coordinate = radial_grid(len(temperature_old))
        diffusion_scale = 1.0 / radius_m**2
        boundary_scale = 1.0 / radius_m
        # Uniform material shrinkage: v_r = Rdot * xi equals the grid velocity.
        # The relative advection term therefore vanishes in material coordinates.
        shrink_velocity = None
    else:
        coordinate = radius_m * radial_grid(len(temperature_old))
        diffusion_scale = 1.0
        boundary_scale = 1.0
        shrink_velocity = None

    for iteration in range(1, max_iterations + 1):
        rho, cp, k, _ = property_model(temperature, moisture)
        heat_storage = rho * cp
        heat_boundary = 2.0 * math.pi * coordinate[-1] * H_T * boundary_scale
        next_temperature = implicit_radial_step(
            old=temperature_old,
            storage=heat_storage,
            conductivity=k,
            dt=dt,
            coordinate=coordinate,
            diffusion_scale=diffusion_scale,
            boundary_conductance=heat_boundary,
            boundary_value=ambient_temperature_c,
            advection_velocity=shrink_velocity,
        )

        _, _, _, diffusivity = property_model(next_temperature, moisture)
        mass_boundary = 2.0 * math.pi * coordinate[-1] * H_M * boundary_scale
        next_moisture = implicit_radial_step(
            old=moisture_old,
            storage=np.ones_like(moisture),
            conductivity=diffusivity,
            dt=dt,
            coordinate=coordinate,
            diffusion_scale=diffusion_scale,
            boundary_conductance=mass_boundary,
            boundary_value=ambient_moisture,
            advection_velocity=shrink_velocity,
            face_conductivity=moisture_faces(next_temperature, moisture, property_model),
        )

        error_t = float(np.max(np.abs(next_temperature - temperature)))
        error_c = float(np.max(np.abs(next_moisture - moisture)))
        temperature = next_temperature
        moisture = next_moisture
        if error_t < tolerance_temperature and error_c < tolerance_moisture:
            break
    else:
        raise RuntimeError(
            f"Picard iteration failed: dT={error_t:.3e}, dC={error_c:.3e}, dt={dt}"
        )

    if not np.all(np.isfinite(temperature)) or not np.all(np.isfinite(moisture)):
        raise FloatingPointError("Non-finite state generated.")
    if np.min(moisture) < -1e-10 or np.max(moisture) > INITIAL_C + 1e-8:
        raise FloatingPointError("Moisture left physical bounds.")
    return temperature, moisture, iteration


def _simulate(environment, property_model, end_time_s, dt_s, output_interval_s,
              stop_at_threshold, radius_history=None, terminal_temp_scale=1.0,
              terminal_moisture_scale=1.0, nodes=None, early_dt_s=1.0,
              terminal_temp_offset_c=0.0):
    """Material-coordinate solve; identical first 3 h stepping for Q2 and Q3.

    Normalized CV volumes are fixed dry-solid mass weights. Their sum is pi.
    Physical dry density scales as R^-2, so total dry mass is constant.
    The reported mass balance is water mass per unit dry-solid mass.
    """
    nodes = NODES if nodes is None else nodes
    coordinate = radial_grid(nodes)
    weights = node_control_volumes(coordinate)[0] / math.pi
    temperature = np.full(nodes, INITIAL_T)
    moisture = np.full(nodes, INITIAL_C)
    times, temperatures, moistures, radii = [0.0], [temperature.copy()], [moisture.copy()], [R_FIXED]
    iterations = []
    threshold_time = math.nan
    threshold_temperature = threshold_moisture = None
    threshold_radius = math.nan
    time_s, next_output, stop_output_time = 0.0, output_interval_s, math.inf
    maximum_time = float(end_time_s if end_time_s is not None else
                         (radius_history.time_s[-1] if radius_history else 10 * 86400))
    monotone_profiles = center_is_max = True
    integrated_loss = maximum_mass_residual = 0.0
    while time_s < maximum_time - 1e-9:
        requested_dt = min(dt_s, early_dt_s) if time_s < 10800 - 1e-9 else dt_s
        step = min(requested_dt, maximum_time - time_s, next_output - time_s)
        for knot in (10800.0, 14400.0):
            if time_s < knot - 1e-9:
                step = min(step, knot - time_s)
        new_time = time_s + step
        old_t, old_c = temperature.copy(), moisture.copy()
        radius = radius_history.values(new_time)[0] if radius_history else R_FIXED
        ambient_t, ambient_c = environment.values(new_time, terminal_temp_scale,
                                                 terminal_moisture_scale)
        if new_time > environment.time_s[-1]:
            ambient_t += terminal_temp_offset_c
        temperature, moisture, count = advance_coupled(
            old_t, old_c, step, radius, 0.0, ambient_t, ambient_c,
            property_model, moving_domain=radius_history is not None)
        iterations.append(count)
        # Backward-Euler surface flux with the same boundary state as the solver.
        boundary_loss = step * 2.0 * H_M / radius * (moisture[-1] - ambient_c)
        residual = float(weights @ (moisture - old_c) + boundary_loss)
        maximum_mass_residual = max(maximum_mass_residual, abs(residual))
        integrated_loss += boundary_loss
        time_s = new_time
        monotone_profiles &= bool(np.all(np.diff(moisture) <= 2e-9))
        center_is_max &= bool(moisture[0] >= np.max(moisture) - 1e-9)
        before, after = float(np.max(old_c)), float(np.max(moisture))
        if before > THRESHOLD_C >= after and math.isnan(threshold_time):
            fraction = (before - THRESHOLD_C) / (before - after)
            threshold_time = time_s - step + fraction * step
            threshold_temperature = old_t + fraction * (temperature - old_t)
            threshold_moisture = old_c + fraction * (moisture - old_c)
            threshold_radius = radius_history.values(threshold_time)[0] if radius_history else R_FIXED
            stop_output_time = math.ceil(threshold_time / output_interval_s) * output_interval_s
        if time_s >= next_output - 1e-8:
            times.append(time_s)
            temperatures.append(temperature.copy())
            moistures.append(moisture.copy())
            radii.append(radius)
            next_output += output_interval_s
        if stop_at_threshold and time_s >= stop_output_time - 1e-8:
            break
    if stop_at_threshold and not math.isfinite(threshold_time):
        raise RuntimeError('Threshold was not reached within supplied time horizon')
    result = {
        'times_s': np.asarray(times), 'temperature_c': np.asarray(temperatures),
        'moisture': np.asarray(moistures), 'threshold_time_s': threshold_time,
        'threshold_temperature_c': threshold_temperature,
        'threshold_moisture': threshold_moisture, 'last_time_s': time_s,
        'picard_max': int(max(iterations, default=0)),
        'picard_mean': float(np.mean(iterations)),
        'monotone_profiles': bool(monotone_profiles), 'center_is_max': bool(center_is_max),
        'maximum_step_mass_residual': maximum_mass_residual,
        'relative_cumulative_mass_residual': abs(float(weights @ moisture) + integrated_loss - INITIAL_C) / INITIAL_C,
        'nodes': nodes, 'late_dt_s': dt_s, 'early_dt_s': min(early_dt_s, dt_s),
        'grid_power': GRID_POWER,
        'transport_frame': 'uniform_material_shrinkage' if radius_history else 'fixed',
    }
    if radius_history:
        result.update(xi=coordinate, radii_m=np.asarray(radii),
                      threshold_radius_m=threshold_radius)
    else:
        result['radius_m'] = coordinate * R_FIXED
    return result


def simulate_fixed(environment, property_model, end_time_s, dt_s, output_interval_s,
                   stop_at_threshold, terminal_temp_scale=1.0,
                   terminal_moisture_scale=1.0, nodes=None, early_dt_s=1.0,
                   terminal_temp_offset_c=0.0):
    return _simulate(environment, property_model, end_time_s, dt_s, output_interval_s,
                     stop_at_threshold, terminal_temp_scale=terminal_temp_scale,
                     terminal_moisture_scale=terminal_moisture_scale, nodes=nodes,
                     early_dt_s=early_dt_s, terminal_temp_offset_c=terminal_temp_offset_c)


def simulate_moving(environment, radius_history, dt_s, output_interval_s,
                    stop_at_threshold=True, terminal_temp_scale=1.0,
                    terminal_moisture_scale=1.0, nodes=None, early_dt_s=1.0,
                    terminal_temp_offset_c=0.0):
    return _simulate(environment, property_q4, None, dt_s, output_interval_s,
                     stop_at_threshold, radius_history=radius_history,
                     terminal_temp_scale=terminal_temp_scale,
                     terminal_moisture_scale=terminal_moisture_scale, nodes=nodes,
                     early_dt_s=early_dt_s, terminal_temp_offset_c=terminal_temp_offset_c)


def extract_fixed_table(
    simulation: dict[str, np.ndarray | float | int | bool],
    requested_times_s: list[float],
    requested_radii_cm: list[float],
    variable: str,
) -> list[list[float]]:
    times = np.asarray(simulation["times_s"])
    radii_m = np.asarray(simulation["radius_m"])
    values = np.asarray(simulation[variable])
    table: list[list[float]] = []
    for requested_time in requested_times_s:
        threshold_time = float(simulation.get("threshold_time_s", math.nan))
        threshold_key = (
            "threshold_temperature_c" if variable == "temperature_c" else "threshold_moisture"
        )
        threshold_profile = simulation.get(threshold_key)
        if (
            threshold_profile is not None
            and math.isfinite(threshold_time)
            and abs(requested_time - threshold_time) < 1e-5
        ):
            profile = np.asarray(threshold_profile)
        else:
            time_index = int(np.argmin(np.abs(times - requested_time)))
            profile = values[time_index]
        row = [
            float(np.interp(radius_cm / 100.0, radii_m, profile))
            for radius_cm in requested_radii_cm
        ]
        table.append(row)
    return table


def extract_moving_profile(
    simulation: dict[str, np.ndarray | float | int | bool],
    requested_time_s: float,
    requested_radii_cm: list[float],
) -> list[float | None]:
    times = np.asarray(simulation["times_s"])
    xi = np.asarray(simulation["xi"])
    radii_m = np.asarray(simulation["radii_m"])
    moisture = np.asarray(simulation["moisture"])
    threshold_time = float(simulation.get("threshold_time_s", math.nan))
    threshold_profile = simulation.get("threshold_moisture")
    if (
        threshold_profile is not None
        and math.isfinite(threshold_time)
        and abs(requested_time_s - threshold_time) < 1e-5
    ):
        profile = np.asarray(threshold_profile)
        radius = float(simulation["threshold_radius_m"])
    else:
        index = int(np.argmin(np.abs(times - requested_time_s)))
        profile = moisture[index]
        radius = float(radii_m[index])
    values: list[float | None] = []
    for radius_cm in requested_radii_cm:
        position_m = radius_cm / 100.0
        if position_m > radius + 1e-12:
            values.append(None)
        else:
            values.append(float(np.interp(position_m / radius, xi, profile)))
    values.append(float(profile[-1]))
    return values


def save_npz(path: Path, simulation: dict[str, np.ndarray | float | int | bool]) -> None:
    arrays = {key: value for key, value in simulation.items() if isinstance(value, np.ndarray)}
    np.savez_compressed(path, **arrays)


def json_ready(value):
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, (np.floating, np.integer)):
        return value.item()
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, dict):
        return {key: json_ready(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_ready(item) for item in value]
    return value


def write_summary(path: Path, payload: dict) -> None:
    path.write_text(
        json.dumps(json_ready(payload), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def run_all(run_convergence: bool = True) -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    for name in ("q1", "q2", "q3", "q4"):
        (RESULTS_DIR / name).mkdir(parents=True, exist_ok=True)

    environment = load_environment()
    radius_history = load_radius_history()
    data_audit = {
        "environment_rows": len(environment.time_s),
        "environment_time_range_s": [environment.time_s[0], environment.time_s[-1]],
        "environment_sampling_s": np.unique(np.diff(environment.time_s)),
        "temperature_range_c": [
            float(np.min(environment.temperature_c)),
            float(np.max(environment.temperature_c)),
        ],
        "ambient_moisture_range": [
            float(np.min(environment.moisture)),
            float(np.max(environment.moisture)),
        ],
        "radius_rows": len(radius_history.time_s),
        "radius_time_range_s": [radius_history.time_s[0], radius_history.time_s[-1]],
        "radius_range_cm": [
            float(np.min(radius_history.radius_m) * 100.0),
            float(np.max(radius_history.radius_m) * 100.0),
        ],
        "radius_monotone_adjustments": radius_history.monotone_adjustments,
    }
    write_summary(RESULTS_DIR / "data_audit.json", data_audit)

    q1 = simulate_fixed(
        environment,
        property_q1,
        end_time_s=1800.0,
        dt_s=1.0,
        output_interval_s=1.0,
        stop_at_threshold=False,
    )
    save_npz(RESULTS_DIR / "q1" / "fields.npz", q1)
    q1_times = [100, 300, 600, 900, 1200, 1500, 1800]
    table_radii = [0.0, 0.5, 1.0, 1.5, 2.0]
    q1_summary = {
        "table_times_s": q1_times,
        "table_radii_cm": table_radii,
        "temperature_table_c": extract_fixed_table(q1, q1_times, table_radii, "temperature_c"),
        "moisture_table": extract_fixed_table(q1, q1_times, table_radii, "moisture"),
        "biot_heat": H_T * R_FIXED / 0.36,
        "initial_diffusivity_m2_s": float(property_q1(
            np.array([INITIAL_T]), np.array([INITIAL_C])
        )[3][0]),
        "biot_mass_initial": H_M * R_FIXED / float(property_q1(
            np.array([INITIAL_T]), np.array([INITIAL_C])
        )[3][0]),
        "picard_max": q1["picard_max"],
        "picard_mean": q1["picard_mean"],
        "temperature_bounds_c": [
            float(np.min(q1["temperature_c"])),
            float(np.max(q1["temperature_c"])),
        ],
        "moisture_bounds": [
            float(np.min(q1["moisture"])),
            float(np.max(q1["moisture"])),
        ],
    }
    write_summary(RESULTS_DIR / "q1" / "summary.json", q1_summary)

    q2 = simulate_fixed(
        environment,
        property_q23,
        end_time_s=10800.0,
        dt_s=1.0,
        output_interval_s=1.0,
        stop_at_threshold=False,
    )
    save_npz(RESULTS_DIR / "q2" / "fields.npz", q2)
    q2_times = [1800, 3600, 5400, 7200, 9000, 10800]
    q2_summary = {
        "table_times_h": [time / 3600.0 for time in q2_times],
        "table_radii_cm": table_radii,
        "temperature_table_c": extract_fixed_table(q2, q2_times, table_radii, "temperature_c"),
        "moisture_table": extract_fixed_table(q2, q2_times, table_radii, "moisture"),
        "picard_max": q2["picard_max"],
        "picard_mean": q2["picard_mean"],
        "monotone_profiles": q2["monotone_profiles"],
        "center_is_max": q2["center_is_max"],
    }
    write_summary(RESULTS_DIR / "q2" / "summary.json", q2_summary)

    q3 = simulate_fixed(
        environment,
        property_q23,
        end_time_s=None,
        dt_s=10.0,
        output_interval_s=60.0,
        stop_at_threshold=True,
    )
    save_npz(RESULTS_DIR / "q3" / "fields.npz", q3)
    q3_end = float(q3["threshold_time_s"])
    q3_report_times = list(np.arange(6 * 3600, math.floor(q3_end / (6 * 3600)) * 6 * 3600 + 1, 6 * 3600))
    if not q3_report_times or abs(q3_report_times[-1] - q3_end) > 1e-6:
        q3_report_times.append(q3_end)
    q3_summary = {
        "drying_time_s": q3_end,
        "drying_time_h": q3_end / 3600.0,
        "output_end_s": q3["last_time_s"],
        "table_times_h": [time / 3600.0 for time in q3_report_times],
        "table_radii_cm": table_radii,
        "moisture_table": extract_fixed_table(q3, q3_report_times, table_radii, "moisture"),
        "center_is_max": q3["center_is_max"],
        "monotone_profiles": q3["monotone_profiles"],
        "picard_max": q3["picard_max"],
        "picard_mean": q3["picard_mean"],
    }

    q4 = simulate_moving(
        environment,
        radius_history,
        dt_s=10.0,
        output_interval_s=60.0,
        stop_at_threshold=True,
    )
    save_npz(RESULTS_DIR / "q4" / "fields.npz", q4)
    q4_end = float(q4["threshold_time_s"])
    q4_report_times = list(np.arange(6 * 3600, math.floor(q4_end / (6 * 3600)) * 6 * 3600 + 1, 6 * 3600))
    if not q4_report_times or abs(q4_report_times[-1] - q4_end) > 1e-6:
        q4_report_times.append(q4_end)
    q4_table_radii = [0.0, 0.5, 1.0]
    q4_summary = {
        "drying_time_s": q4_end,
        "drying_time_h": q4_end / 3600.0,
        "output_end_s": q4["last_time_s"],
        "table_times_h": [time / 3600.0 for time in q4_report_times],
        "table_radii_cm": q4_table_radii,
        "columns": q4_table_radii + ["药材表面"],
        "moisture_table": [
            extract_moving_profile(q4, time, q4_table_radii) for time in q4_report_times
        ],
        "radius_at_end_cm": float(radius_history.values(q4_end)[0] * 100.0),
        "center_is_max": q4["center_is_max"],
        "monotone_profiles": q4["monotone_profiles"],
        "picard_max": q4["picard_max"],
        "picard_mean": q4["picard_mean"],
    }

    if run_convergence:
        q3_fine = simulate_fixed(
            environment,
            property_q23,
            end_time_s=None,
            dt_s=5.0,
            output_interval_s=60.0,
            stop_at_threshold=True,
        )
        q4_fine = simulate_moving(
            environment,
            radius_history,
            dt_s=5.0,
            output_interval_s=60.0,
            stop_at_threshold=True,
        )
        q3_summary["time_step_check"] = {
            "dt10_s_time_h": q3_end / 3600.0,
            "dt5_s_time_h": float(q3_fine["threshold_time_s"]) / 3600.0,
            "relative_difference": abs(
                float(q3_fine["threshold_time_s"]) - q3_end
            ) / float(q3_fine["threshold_time_s"]),
        }
        q4_summary["time_step_check"] = {
            "dt10_s_time_h": q4_end / 3600.0,
            "dt5_s_time_h": float(q4_fine["threshold_time_s"]) / 3600.0,
            "relative_difference": abs(
                float(q4_fine["threshold_time_s"]) - q4_end
            ) / float(q4_fine["threshold_time_s"]),
        }

    for name, result in [('q1', q1), ('q2', q2), ('q3', q3), ('q4', q4)]:
        write_summary(RESULTS_DIR / name / 'diagnostics.json', {
            key: value for key, value in result.items() if not isinstance(value, np.ndarray)
            and value is not None})
    write_summary(RESULTS_DIR / "q3" / "summary.json", q3_summary)
    write_summary(RESULTS_DIR / "q4" / "summary.json", q4_summary)

    print(json.dumps(
        {
            "q1_surface_temperature_1800_c": q1_summary["temperature_table_c"][-1][-1],
            "q1_center_moisture_1800": q1_summary["moisture_table"][-1][0],
            "q2_center_temperature_3h_c": q2_summary["temperature_table_c"][-1][0],
            "q2_center_moisture_3h": q2_summary["moisture_table"][-1][0],
            "q3_drying_time_h": q3_summary["drying_time_h"],
            "q4_drying_time_h": q4_summary["drying_time_h"],
            "q4_radius_at_end_cm": q4_summary["radius_at_end_cm"],
        },
        ensure_ascii=False,
        indent=2,
    ))


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--skip-convergence", action="store_true")
    args = parser.parse_args()
    run_all(run_convergence=not args.skip_convergence)


if __name__ == "__main__":
    main()
