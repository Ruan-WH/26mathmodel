# File: code/solve_drying.py
"""Numerical solution for CUMCM 2026 Problem A.

The model uses a node-centred finite-volume discretisation in the radial
coordinate, fully implicit diffusion, Picard iteration for nonlinear
properties, and an ALE-style fixed-domain transform for shrinkage.
"""

from __future__ import annotations

import argparse
import json
import math
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator
from scipy.linalg import solve_banded


ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = ROOT / "附件"
RESULTS_DIR = ROOT / "results"
H_T = 25.0
H_M = 8.0e-7
R_FIXED = 0.02
NODES = 161
OUTPUT_NODES = 21
GRID_POWER = 2.0
INITIAL_T = 28.0
INITIAL_C = 2.55
THRESHOLD_C = 0.15
TERMINAL_TEMPERATURE_C = 50.0
TERMINAL_MOISTURE = 0.05


@dataclass
class Environment:
    time_s: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray
    temp_interp: PchipInterpolator
    moisture_interp: PchipInterpolator

    def values(
        self,
        time_s: float,
        terminal_temp_scale: float = 1.0,
        terminal_moisture_scale: float = 1.0,
    ) -> tuple[float, float]:
        cutoff = float(self.time_s[-1])
        if time_s <= cutoff:
            return float(self.temp_interp(time_s)), float(self.moisture_interp(time_s))
        return (
            float(TERMINAL_TEMPERATURE_C * terminal_temp_scale),
            float(TERMINAL_MOISTURE * terminal_moisture_scale),
        )


@dataclass
class RadiusHistory:
    time_s: np.ndarray
    radius_m: np.ndarray
    interp: PchipInterpolator
    derivative: PchipInterpolator
    monotone_adjustments: int

    def values(self, time_s: float) -> tuple[float, float]:
        bounded = min(max(time_s, float(self.time_s[0])), float(self.time_s[-1]))
        return float(self.interp(bounded)), float(self.derivative(bounded))


def _read_numeric_sheet(path: Path) -> np.ndarray:
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book.active
    rows: list[list[float]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        rows.append([float(value) for value in row])
    return np.asarray(rows, dtype=float)


def load_environment() -> Environment:
    data = _read_numeric_sheet(INPUT_DIR / "附件1.xlsx")
    time_s, temperature_c, moisture = data.T
    if not np.all(np.diff(time_s) > 0):
        raise ValueError("附件1时间列不是严格递增。")
    if not np.all(np.isfinite(data)):
        raise ValueError("附件1含非有限数。")
    return Environment(
        time_s=time_s,
        temperature_c=temperature_c,
        moisture=moisture,
        temp_interp=PchipInterpolator(time_s, temperature_c),
        moisture_interp=PchipInterpolator(time_s, moisture),
    )


def load_radius_history() -> RadiusHistory:
    data = _read_numeric_sheet(INPUT_DIR / "附件2.xlsx")
    time_s = data[:, 0]
    raw_radius_m = data[:, 1] / 100.0
    if not np.all(np.diff(time_s) > 0):
        raise ValueError("附件2时间列不是严格递增。")
    monotone_radius_m = np.minimum.accumulate(raw_radius_m)
    adjustments = int(np.count_nonzero(np.abs(monotone_radius_m - raw_radius_m) > 1e-12))
    interp = PchipInterpolator(time_s, monotone_radius_m, extrapolate=False)
    return RadiusHistory(
        time_s=time_s,
        radius_m=monotone_radius_m,
        interp=interp,
        derivative=interp.derivative(),
        monotone_adjustments=adjustments,
    )


def property_q1(temperature_c: np.ndarray, moisture: np.ndarray) -> tuple[np.ndarray, ...]:
    del temperature_c
    rho = np.full_like(moisture, 820.0)
    cp = np.full_like(moisture, 2600.0)
    k = np.full_like(moisture, 0.36)
    safe_c = np.maximum(moisture, 1e-8)
    diffusivity = 7.0e-9 * np.exp(-0.89 / safe_c)
    return rho, cp, k, diffusivity


def property_q23(temperature_c: np.ndarray, moisture: np.ndarray) -> tuple[np.ndarray, ...]:
    safe_c = np.maximum(moisture, 1e-8)
    rho = 650.0 + 128.0 * safe_c
    cp = 1450.0 + 2736.0 * safe_c / (safe_c + 1.0)
    k = 0.21 + 0.38 * safe_c / (safe_c + 1.0)
    diffusivity = (
        2.4e-3
        * np.exp(-0.45 / safe_c)
        * np.exp(-3850.0 / (temperature_c + 273.15))
    )
    return rho, cp, k, diffusivity


def property_q4(temperature_c: np.ndarray, moisture: np.ndarray) -> tuple[np.ndarray, ...]:
    safe_c = np.maximum(moisture, 1e-8)
    rho = 760.0 + 90.0 * safe_c
    cp = 1850.0 + 2150.0 * safe_c / (safe_c + 1.0)
    k = 0.12 + 0.20 * safe_c / (safe_c + 1.0)
    diffusivity = (
        4.2e-4
        * np.exp(-0.30 / safe_c)
        * np.exp(-3850.0 / (temperature_c + 273.15))
    )
    return rho, cp, k, diffusivity


def face_interpolate(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    """Linear face interpolation for a smooth state-dependent coefficient."""
    return 0.5 * (left + right)


def radial_grid(count: int, radius: float = 1.0) -> np.ndarray:
    """Surface-refined computational mesh; output positions are independent."""
    s = np.linspace(0.0, 1.0, count)
    return radius * (1.0 - (1.0 - s) ** GRID_POWER)


def node_control_volumes(coordinate: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    """Return dimensionless/physical radial CV volumes and internal face radii."""
    faces = 0.5 * (coordinate[:-1] + coordinate[1:])
    lower = np.concatenate(([0.0], faces))
    upper = np.concatenate((faces, [coordinate[-1]]))
    volumes = math.pi * (upper**2 - lower**2)
    return volumes, faces


def implicit_radial_step(
    old: np.ndarray,
    storage: np.ndarray,
    conductivity: np.ndarray,
    dt: float,
    coordinate: np.ndarray,
    diffusion_scale: float,
    boundary_conductance: float,
    boundary_value: float,
    advection_velocity: np.ndarray | None = None,
) -> np.ndarray:
    """One implicit node-centred radial FVM step on coordinate [0, 1] or [0, R]."""
    count = len(old)
    volumes, faces = node_control_volumes(coordinate)
    delta = np.diff(coordinate)
    face_property = face_interpolate(conductivity[:-1], conductivity[1:])
    conductance = 2.0 * math.pi * faces * face_property * diffusion_scale / delta

    capacity = storage * volumes / dt
    diagonal = capacity.copy()
    lower = np.zeros(count - 1)
    upper = np.zeros(count - 1)
    rhs = capacity * old

    diagonal[:-1] += conductance
    diagonal[1:] += conductance
    upper -= conductance
    lower -= conductance

    diagonal[-1] += boundary_conductance
    rhs[-1] += boundary_conductance * boundary_value

    if advection_velocity is not None:
        # PDE form: phi_t + u(xi) phi_xi = diffusion; u >= 0 under shrinkage.
        advective = storage[1:] * volumes[1:] * np.maximum(advection_velocity[1:], 0.0) / delta
        diagonal[1:] += advective
        lower -= advective

    matrix = np.zeros((3, count))
    matrix[0, 1:] = upper
    matrix[1, :] = diagonal
    matrix[2, :-1] = lower
    return solve_banded((1, 1), matrix, rhs, check_finite=False)


def advance_coupled(
    temperature_old: np.ndarray,
    moisture_old: np.ndarray,
    dt: float,
    radius_m: float,
    radius_rate_m_s: float,
    ambient_temperature_c: float,
    ambient_moisture: float,
    property_model: Callable[[np.ndarray, np.ndarray], tuple[np.ndarray, ...]],
    moving_domain: bool,
    tolerance_temperature: float = 1e-8,
    tolerance_moisture: float = 1e-10,
    max_iterations: int = 30,
    balance: dict | None = None,
) -> tuple[np.ndarray, np.ndarray, int]:
    temperature = temperature_old.copy()
    moisture = moisture_old.copy()
    if moving_domain:
        coordinate = radial_grid(len(temperature_old))
        diffusion_scale = 1.0 / radius_m**2
        boundary_scale = 1.0 / radius_m
        shrink_velocity = -(radius_rate_m_s / radius_m) * coordinate
    else:
        coordinate = radial_grid(len(temperature_old), radius_m)
        diffusion_scale = 1.0
        boundary_scale = 1.0
        shrink_velocity = None

    for iteration in range(1, max_iterations + 1):
        rho, cp, k, _ = property_model(temperature, moisture)
        heat_storage = rho * cp
        heat_boundary = 2.0 * math.pi * coordinate[-1] * H_T * boundary_scale
        next_temperature = implicit_radial_step(
            old=temperature_old,
            storage=heat_storage,
            conductivity=k,
            dt=dt,
            coordinate=coordinate,
            diffusion_scale=diffusion_scale,
            boundary_conductance=heat_boundary,
            boundary_value=ambient_temperature_c,
            advection_velocity=shrink_velocity,
        )

        _, _, _, diffusivity = property_model(next_temperature, moisture)
        mass_boundary = 2.0 * math.pi * coordinate[-1] * H_M * boundary_scale
        next_moisture = implicit_radial_step(
            old=moisture_old,
            storage=np.ones_like(moisture),
            conductivity=diffusivity,
            dt=dt,
            coordinate=coordinate,
            diffusion_scale=diffusion_scale,
            boundary_conductance=mass_boundary,
            boundary_value=ambient_moisture,
            advection_velocity=shrink_velocity,
        )

        error_t = float(np.max(np.abs(next_temperature - temperature)))
        error_c = float(np.max(np.abs(next_moisture - moisture)))
        temperature = next_temperature
        moisture = next_moisture
        if error_t < tolerance_temperature and error_c < tolerance_moisture:
            break
    else:
        raise RuntimeError(
            f"Picard iteration failed: dT={error_t:.3e}, dC={error_c:.3e}, dt={dt}"
        )

    if not np.all(np.isfinite(temperature)) or not np.all(np.isfinite(moisture)):
        raise FloatingPointError("Non-finite state generated.")
    if np.min(moisture) < -1e-10 or np.max(moisture) > INITIAL_C + 1e-8:
        raise FloatingPointError("Moisture left physical bounds.")
    if balance is not None:
        volumes, _ = node_control_volumes(coordinate)
        adv_c = adv_t = 0.0
        if shrink_velocity is not None:
            weight = volumes[1:] * np.maximum(shrink_velocity[1:], 0.0) / np.diff(coordinate)
            adv_c = dt * float(np.dot(weight, np.diff(moisture)))
            adv_t = dt * float(np.dot(weight * heat_storage[1:], np.diff(temperature)))
        change = float(np.dot(volumes, moisture - moisture_old))
        flux = dt * mass_boundary * (float(moisture[-1]) - ambient_moisture)
        reference = INITIAL_C * float(volumes.sum())
        balance['moisture_change'] += change / reference
        balance['boundary_loss'] += flux / reference
        balance['relative_transport'] += adv_c / reference
        residual = change + flux + adv_c
        balance['maximum_step_mass_residual'] = max(balance['maximum_step_mass_residual'], abs(residual) / reference)
        heat_change = float(np.dot(heat_storage * volumes, temperature - temperature_old))
        heat_flux = dt * heat_boundary * (float(temperature[-1]) - ambient_temperature_c)
        heat_reference = INITIAL_T * float(np.dot(heat_storage, volumes))
        heat_residual = (heat_change + heat_flux + adv_t) / heat_reference
        balance['maximum_step_heat_equation_residual'] = max(balance['maximum_step_heat_equation_residual'], abs(heat_residual))
    return temperature, moisture, iteration


def new_balance() -> dict:
    return dict(moisture_change=0.0, boundary_loss=0.0, relative_transport=0.0,
                maximum_step_mass_residual=0.0, maximum_step_heat_equation_residual=0.0)


def simulate_fixed(
    environment: Environment,
    property_model: Callable[[np.ndarray, np.ndarray], tuple[np.ndarray, ...]],
    end_time_s: float | None,
    dt_s: float,
    output_interval_s: float,
    stop_at_threshold: bool,
    terminal_temp_scale: float = 1.0,
    terminal_moisture_scale: float = 1.0,
    stop_interval_s: float | None = None,
) -> dict[str, np.ndarray | float | int | bool]:
    temperature = np.full(NODES, INITIAL_T)
    moisture = np.full(NODES, INITIAL_C)
    radial_m = radial_grid(NODES, R_FIXED)
    output_r = np.linspace(0.0, R_FIXED, OUTPUT_NODES)
    times = [0.0]
    temperatures = [np.interp(output_r, radial_m, temperature)]
    moistures = [np.interp(output_r, radial_m, moisture)]
    iterations: list[int] = []
    monotone_profiles = True
    center_is_max = True
    previous_max = float(np.max(moisture))
    previous_time = 0.0
    threshold_time = math.nan
    threshold_temperature: np.ndarray | None = None
    threshold_moisture: np.ndarray | None = None
    stop_output_time = math.inf
    balance = new_balance()
    time_s = 0.0
    next_output = output_interval_s
    maximum_time = float(end_time_s if end_time_s is not None else 10.0 * 24.0 * 3600.0)

    while time_s < maximum_time - 1e-9:
        step = min(dt_s, maximum_time - time_s, next_output - time_s)
        new_time = time_s + step
        temperature_before = temperature.copy()
        moisture_before = moisture.copy()
        ambient_t, ambient_c = environment.values(
            new_time, terminal_temp_scale, terminal_moisture_scale
        )
        temperature, moisture, count = advance_coupled(
            temperature,
            moisture,
            step,
            R_FIXED,
            0.0,
            ambient_t,
            ambient_c,
            property_model,
            moving_domain=False,
            balance=balance,
        )
        iterations.append(count)
        time_s = new_time
        current_max = float(np.max(moisture))
        monotone_profiles = monotone_profiles and bool(np.all(np.diff(moisture) <= 2e-9))
        center_is_max = center_is_max and bool(
            moisture[0] >= float(np.max(moisture)) - 1e-9
        )

        if previous_max > THRESHOLD_C >= current_max and math.isnan(threshold_time):
            fraction = (previous_max - THRESHOLD_C) / max(previous_max - current_max, 1e-15)
            threshold_time = previous_time + fraction * step
            threshold_temperature = (
                temperature_before + fraction * (temperature - temperature_before)
            )
            threshold_moisture = moisture_before + fraction * (moisture - moisture_before)
            stop_spacing = output_interval_s if stop_interval_s is None else stop_interval_s
            stop_output_time = math.ceil(threshold_time / stop_spacing) * stop_spacing

        if time_s + 1e-8 >= next_output:
            times.append(time_s)
            temperatures.append(np.interp(output_r, radial_m, temperature))
            moistures.append(np.interp(output_r, radial_m, moisture))
            next_output += output_interval_s

        if stop_at_threshold and time_s + 1e-8 >= stop_output_time:
            break
        previous_max = current_max
        previous_time = time_s

    return {
        "times_s": np.asarray(times),
        "radius_m": output_r,
        "temperature_c": np.asarray(temperatures),
        "moisture": np.asarray(moistures),
        "threshold_time_s": threshold_time,
        "threshold_temperature_c": None if threshold_temperature is None else np.interp(output_r, radial_m, threshold_temperature),
        "threshold_moisture": None if threshold_moisture is None else np.interp(output_r, radial_m, threshold_moisture),
        "last_time_s": time_s,
        "picard_max": int(max(iterations, default=0)),
        "picard_mean": float(np.mean(iterations) if iterations else 0.0),
        "monotone_profiles": monotone_profiles,
        "center_is_max": center_is_max,
        "balance": balance,
    }


def simulate_moving(
    environment: Environment,
    radius_history: RadiusHistory,
    dt_s: float,
    output_interval_s: float,
    stop_at_threshold: bool = True,
    terminal_temp_scale: float = 1.0,
    terminal_moisture_scale: float = 1.0,
    include_advection: bool = False,
) -> dict[str, np.ndarray | float | int | bool]:
    # Default: homogeneous solid shrinkage, material velocity equals mesh velocity.
    # True retains the former stationary-spatial scalar model for structural comparison.
    xi = radial_grid(NODES)
    # Preserve native nodes; interpolate to physical positions only at export.
    output_xi = xi.copy()
    temperature = np.full(NODES, INITIAL_T)
    moisture = np.full(NODES, INITIAL_C)
    times = [0.0]
    radii = [R_FIXED]
    temperatures = [np.interp(output_xi, xi, temperature)]
    moistures = [np.interp(output_xi, xi, moisture)]
    iterations: list[int] = []
    previous_max = float(np.max(moisture))
    previous_time = 0.0
    threshold_time = math.nan
    threshold_temperature: np.ndarray | None = None
    threshold_moisture: np.ndarray | None = None
    threshold_radius_m = math.nan
    stop_output_time = math.inf
    time_s = 0.0
    next_output = output_interval_s
    maximum_time = float(radius_history.time_s[-1])
    monotone_profiles = True
    center_is_max = True
    maximum_cfl = 0.0
    balance = new_balance()

    while time_s < maximum_time - 1e-9:
        step = min(dt_s, maximum_time - time_s, next_output - time_s)
        new_time = time_s + step
        temperature_before = temperature.copy()
        moisture_before = moisture.copy()
        radius_m, radius_rate_m_s = radius_history.values(new_time)
        if not include_advection:
            radius_rate_m_s = 0.0
        ambient_t, ambient_c = environment.values(
            new_time, terminal_temp_scale, terminal_moisture_scale
        )
        temperature, moisture, count = advance_coupled(
            temperature,
            moisture,
            step,
            radius_m,
            radius_rate_m_s,
            ambient_t,
            ambient_c,
            property_q4,
            moving_domain=True,
            balance=balance,
        )
        iterations.append(count)
        time_s = new_time
        current_max = float(np.max(moisture))
        maximum_cfl = max(
            maximum_cfl,
            float(np.max(abs(radius_rate_m_s / radius_m) * xi[1:] * step / np.diff(xi))),
        )
        monotone_profiles = monotone_profiles and bool(np.all(np.diff(moisture) <= 2e-9))
        center_is_max = center_is_max and bool(
            moisture[0] >= float(np.max(moisture)) - 1e-9
        )

        if previous_max > THRESHOLD_C >= current_max and math.isnan(threshold_time):
            fraction = (previous_max - THRESHOLD_C) / max(previous_max - current_max, 1e-15)
            threshold_time = previous_time + fraction * step
            threshold_temperature = (
                temperature_before + fraction * (temperature - temperature_before)
            )
            threshold_moisture = moisture_before + fraction * (moisture - moisture_before)
            threshold_radius_m = radius_history.values(threshold_time)[0]
            stop_output_time = (
                math.ceil(threshold_time / output_interval_s) * output_interval_s
            )

        if time_s + 1e-8 >= next_output:
            times.append(time_s)
            radii.append(radius_m)
            temperatures.append(np.interp(output_xi, xi, temperature))
            moistures.append(np.interp(output_xi, xi, moisture))
            next_output += output_interval_s

        if stop_at_threshold and time_s + 1e-8 >= stop_output_time:
            break
        previous_max = current_max
        previous_time = time_s

    return {
        "times_s": np.asarray(times),
        "xi": output_xi,
        "radii_m": np.asarray(radii),
        "temperature_c": np.asarray(temperatures),
        "moisture": np.asarray(moistures),
        "threshold_time_s": threshold_time,
        "threshold_temperature_c": None if threshold_temperature is None else np.interp(output_xi, xi, threshold_temperature),
        "threshold_moisture": None if threshold_moisture is None else np.interp(output_xi, xi, threshold_moisture),
        "threshold_radius_m": threshold_radius_m,
        "last_time_s": time_s,
        "picard_max": int(max(iterations, default=0)),
        "picard_mean": float(np.mean(iterations) if iterations else 0.0),
        "monotone_profiles": monotone_profiles,
        "center_is_max": center_is_max,
        "maximum_advection_cfl": maximum_cfl,
        "balance": balance,
        "velocity_model": "stationary_spatial_comparison" if include_advection else "uniform_material_shrinkage",
    }


def extract_fixed_table(
    simulation: dict[str, np.ndarray | float | int | bool],
    requested_times_s: list[float],
    requested_radii_cm: list[float],
    variable: str,
) -> list[list[float]]:
    times = np.asarray(simulation["times_s"])
    radii_m = np.asarray(simulation["radius_m"])
    values = np.asarray(simulation[variable])
    table: list[list[float]] = []
    for requested_time in requested_times_s:
        threshold_time = float(simulation.get("threshold_time_s", math.nan))
        threshold_key = (
            "threshold_temperature_c" if variable == "temperature_c" else "threshold_moisture"
        )
        threshold_profile = simulation.get(threshold_key)
        if (
            threshold_profile is not None
            and math.isfinite(threshold_time)
            and abs(requested_time - threshold_time) < 1e-5
        ):
            profile = np.asarray(threshold_profile)
        else:
            time_index = int(np.argmin(np.abs(times - requested_time)))
            profile = values[time_index]
        row = [
            float(np.interp(radius_cm / 100.0, radii_m, profile))
            for radius_cm in requested_radii_cm
        ]
        table.append(row)
    return table


def extract_moving_profile(
    simulation: dict[str, np.ndarray | float | int | bool],
    requested_time_s: float,
    requested_radii_cm: list[float],
) -> list[float | None]:
    times = np.asarray(simulation["times_s"])
    xi = np.asarray(simulation["xi"])
    radii_m = np.asarray(simulation["radii_m"])
    moisture = np.asarray(simulation["moisture"])
    threshold_time = float(simulation.get("threshold_time_s", math.nan))
    threshold_profile = simulation.get("threshold_moisture")
    if (
        threshold_profile is not None
        and math.isfinite(threshold_time)
        and abs(requested_time_s - threshold_time) < 1e-5
    ):
        profile = np.asarray(threshold_profile)
        radius = float(simulation["threshold_radius_m"])
    else:
        index = int(np.argmin(np.abs(times - requested_time_s)))
        profile = moisture[index]
        radius = float(radii_m[index])
    values: list[float | None] = []
    for radius_cm in requested_radii_cm:
        position_m = radius_cm / 100.0
        if position_m > radius + 1e-12:
            values.append(None)
        else:
            values.append(float(np.interp(position_m / radius, xi, profile)))
    values.append(float(profile[-1]))
    return values


def save_npz(path: Path, simulation: dict[str, np.ndarray | float | int | bool]) -> None:
    arrays = {key: value for key, value in simulation.items() if isinstance(value, np.ndarray)}
    np.savez_compressed(path, **arrays)


def json_ready(value):
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, (np.floating, np.integer)):
        return value.item()
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, dict):
        return {key: json_ready(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_ready(item) for item in value]
    return value


def write_summary(path: Path, payload: dict) -> None:
    path.write_text(
        json.dumps(json_ready(payload), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def run_all(run_convergence: bool = True) -> None:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    for name in ("q1", "q2", "q3", "q4"):
        (RESULTS_DIR / name).mkdir(parents=True, exist_ok=True)

    environment = load_environment()
    radius_history = load_radius_history()
    data_audit = {
        "environment_rows": len(environment.time_s),
        "environment_time_range_s": [environment.time_s[0], environment.time_s[-1]],
        "environment_sampling_s": np.unique(np.diff(environment.time_s)),
        "temperature_range_c": [
            float(np.min(environment.temperature_c)),
            float(np.max(environment.temperature_c)),
        ],
        "ambient_moisture_range": [
            float(np.min(environment.moisture)),
            float(np.max(environment.moisture)),
        ],
        "terminal_boundary_after_14400_s": {
            "temperature_c": TERMINAL_TEMPERATURE_C,
            "moisture": TERMINAL_MOISTURE,
            "basis": "constant-stage setpoints; final-hour attachment means are 49.9989 C and 0.049988",
        },
        "radius_rows": len(radius_history.time_s),
        "radius_time_range_s": [radius_history.time_s[0], radius_history.time_s[-1]],
        "radius_range_cm": [
            float(np.min(radius_history.radius_m) * 100.0),
            float(np.max(radius_history.radius_m) * 100.0),
        ],
        "radius_monotone_adjustments": radius_history.monotone_adjustments,
    }
    write_summary(RESULTS_DIR / "data_audit.json", data_audit)

    q1 = simulate_fixed(
        environment,
        property_q1,
        end_time_s=1800.0,
        dt_s=1.0,
        output_interval_s=1.0,
        stop_at_threshold=False,
    )
    save_npz(RESULTS_DIR / "q1" / "fields.npz", q1)
    q1_times = [100, 300, 600, 900, 1200, 1500, 1800]
    table_radii = [0.0, 0.5, 1.0, 1.5, 2.0]
    q1_summary = {
        "table_times_s": q1_times,
        "table_radii_cm": table_radii,
        "temperature_table_c": extract_fixed_table(q1, q1_times, table_radii, "temperature_c"),
        "moisture_table": extract_fixed_table(q1, q1_times, table_radii, "moisture"),
        "biot_heat": H_T * R_FIXED / 0.36,
        "initial_diffusivity_m2_s": float(property_q1(
            np.array([INITIAL_T]), np.array([INITIAL_C])
        )[3][0]),
        "biot_mass_initial": H_M * R_FIXED / float(property_q1(
            np.array([INITIAL_T]), np.array([INITIAL_C])
        )[3][0]),
        "picard_max": q1["picard_max"],
        "picard_mean": q1["picard_mean"],
        "temperature_bounds_c": [
            float(np.min(q1["temperature_c"])),
            float(np.max(q1["temperature_c"])),
        ],
        "moisture_bounds": [
            float(np.min(q1["moisture"])),
            float(np.max(q1["moisture"])),
        ],
    }
    write_summary(RESULTS_DIR / "q1" / "summary.json", q1_summary)

    q2 = simulate_fixed(
        environment,
        property_q23,
        end_time_s=None,
        dt_s=1.0,
        output_interval_s=1.0,
        stop_at_threshold=True,
        stop_interval_s=60.0,
    )
    save_npz(RESULTS_DIR / "q2" / "fields.npz", q2)
    q2_times = [1800, 3600, 5400, 7200, 9000, 10800]
    q2_summary = {
        "table_times_h": [time / 3600.0 for time in q2_times],
        "table_radii_cm": table_radii,
        "temperature_table_c": extract_fixed_table(q2, q2_times, table_radii, "temperature_c"),
        "moisture_table": extract_fixed_table(q2, q2_times, table_radii, "moisture"),
        "picard_max": q2["picard_max"],
        "picard_mean": q2["picard_mean"],
        "monotone_profiles": q2["monotone_profiles"],
        "center_is_max": q2["center_is_max"],
    }
    write_summary(RESULTS_DIR / "q2" / "summary.json", q2_summary)

    q3 = dict(q2)
    for key in ("times_s", "temperature_c", "moisture"):
        q3[key] = q2[key][::60].copy()
    save_npz(RESULTS_DIR / "q3" / "fields.npz", q3)
    q3_end = float(q3["threshold_time_s"])
    q3_report_times = list(np.arange(6 * 3600, math.floor(q3_end / (6 * 3600)) * 6 * 3600 + 1, 6 * 3600))
    if not q3_report_times or abs(q3_report_times[-1] - q3_end) > 1e-6:
        q3_report_times.append(q3_end)
    q3_summary = {
        "drying_time_s": q3_end,
        "drying_time_h": q3_end / 3600.0,
        "output_end_s": q3["last_time_s"],
        "table_times_h": [time / 3600.0 for time in q3_report_times],
        "table_radii_cm": table_radii,
        "moisture_table": extract_fixed_table(q3, q3_report_times, table_radii, "moisture"),
        "center_is_max": q3["center_is_max"],
        "monotone_profiles": q3["monotone_profiles"],
        "picard_max": q3["picard_max"],
        "picard_mean": q3["picard_mean"],
    }

    q4 = simulate_moving(
        environment,
        radius_history,
        dt_s=1.0,
        output_interval_s=60.0,
        stop_at_threshold=True,
    )
    save_npz(RESULTS_DIR / "q4" / "fields.npz", q4)
    q4_end = float(q4["threshold_time_s"])
    q4_report_times = list(np.arange(6 * 3600, math.floor(q4_end / (6 * 3600)) * 6 * 3600 + 1, 6 * 3600))
    if not q4_report_times or abs(q4_report_times[-1] - q4_end) > 1e-6:
        q4_report_times.append(q4_end)
    q4_table_radii = [0.0, 0.5, 1.0]
    q4_summary = {
        "drying_time_s": q4_end,
        "drying_time_h": q4_end / 3600.0,
        "output_end_s": q4["last_time_s"],
        "table_times_h": [time / 3600.0 for time in q4_report_times],
        "table_radii_cm": q4_table_radii,
        "columns": q4_table_radii + ["药材表面"],
        "moisture_table": [
            extract_moving_profile(q4, time, q4_table_radii) for time in q4_report_times
        ],
        "radius_at_end_cm": float(radius_history.values(q4_end)[0] * 100.0),
        "center_is_max": q4["center_is_max"],
        "monotone_profiles": q4["monotone_profiles"],
        "maximum_advection_cfl": q4["maximum_advection_cfl"],
        "picard_max": q4["picard_max"],
        "picard_mean": q4["picard_mean"],
    }

    if run_convergence:
        q3_dt2 = simulate_fixed(
            environment,
            property_q23,
            end_time_s=None,
            dt_s=2.0,
            output_interval_s=60.0,
            stop_at_threshold=True,
        )
        q4_dt2 = simulate_moving(
            environment,
            radius_history,
            dt_s=2.0,
            output_interval_s=60.0,
            stop_at_threshold=True,
        )
        q4_stationary = simulate_moving(
            environment,
            radius_history,
            dt_s=1.0,
            output_interval_s=60.0,
            stop_at_threshold=True,
            include_advection=True,
        )
        q3_summary["time_step_check"] = {
            "dt1_s_time_h": q3_end / 3600.0,
            "dt2_s_time_h": float(q3_dt2["threshold_time_s"]) / 3600.0,
            "relative_difference": abs(
                float(q3_dt2["threshold_time_s"]) - q3_end
            ) / float(q3_dt2["threshold_time_s"]),
        }
        q4_summary["time_step_check"] = {
            "dt1_s_time_h": q4_end / 3600.0,
            "dt2_s_time_h": float(q4_dt2["threshold_time_s"]) / 3600.0,
            "relative_difference": abs(
                float(q4_dt2["threshold_time_s"]) - q4_end
            ) / float(q4_dt2["threshold_time_s"]),
        }
        q4_summary["moving_term_comparison"] = {
            "material_model_h": q4_end / 3600.0,
            "stationary_spatial_model_h": float(q4_stationary["threshold_time_s"]) / 3600.0,
            "difference_h": (
                float(q4_stationary["threshold_time_s"]) - q4_end
            ) / 3600.0,
        }

    for summary, simulation in [(q1_summary, q1), (q2_summary, q2), (q3_summary, q3), (q4_summary, q4)]:
        summary['numerics'] = {'nodes': NODES, 'grid_power': GRID_POWER, 'dt_s': 1.0, 'output_nodes': simulation['moisture'].shape[1]}
        summary['balance'] = simulation['balance']
        summary['balance']['cumulative_mass_equation_residual'] = abs(sum(simulation['balance'][key] for key in ['moisture_change','boundary_loss','relative_transport']))
    q2_summary['output_end_s'] = q2['last_time_s']
    q2_summary['drying_time_h'] = q2['threshold_time_s'] / 3600.0
    q3_summary['shared_trajectory_with_q2'] = bool(np.array_equal(q3['moisture'], q2['moisture'][::60]))
    q4_summary['velocity_model'] = q4['velocity_model']
    write_summary(RESULTS_DIR / "q1" / "summary.json", q1_summary)
    write_summary(RESULTS_DIR / "q2" / "summary.json", q2_summary)
    write_summary(RESULTS_DIR / "q3" / "summary.json", q3_summary)
    write_summary(RESULTS_DIR / "q4" / "summary.json", q4_summary)

    print(json.dumps(
        {
            "q1_surface_temperature_1800_c": q1_summary["temperature_table_c"][-1][-1],
            "q1_center_moisture_1800": q1_summary["moisture_table"][-1][0],
            "q2_center_temperature_3h_c": q2_summary["temperature_table_c"][-1][0],
            "q2_center_moisture_3h": q2_summary["moisture_table"][-1][0],
            "q3_drying_time_h": q3_summary["drying_time_h"],
            "q4_drying_time_h": q4_summary["drying_time_h"],
            "q4_radius_at_end_cm": q4_summary["radius_at_end_cm"],
        },
        ensure_ascii=False,
        indent=2,
    ))


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--skip-convergence", action="store_true")
    args = parser.parse_args()
    run_all(run_convergence=not args.skip_convergence)


if __name__ == "__main__":
    main()
