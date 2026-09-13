# File: code/write_results.py
"""Write the four official-format Excel result workbooks."""

from __future__ import annotations

import json
from copy import copy
import hashlib
import shutil
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "附件" / "附件3"
RESULTS_DIR = ROOT / "results"
def reset_sheet(sheet) -> dict:
    """Clear template placeholders while retaining the official sheet settings."""
    last_column = sheet.max_column
    template = {
        "corner": sheet.cell(1, 1).value,
        "styles": {(row, column): copy(sheet.cell(row, column)._style)
                   for row in (1, 2) for column in range(1, last_column + 1)},
        "last_column": last_column,
        "column_a": copy(sheet.column_dimensions["A"]),
        "column_data": copy(sheet.column_dimensions["B"]),
        "row_header": copy(sheet.row_dimensions[1]),
        "row_data": copy(sheet.row_dimensions[2]),
    }
    for row in sheet.iter_rows():
        for cell in row:
            cell.value = None
    return template


def style_sheet(sheet, row_count: int, column_count: int, template: dict) -> None:
    """Extend template styles; only result cells use the required four decimals."""
    last_template_column = template["last_column"]
    for row in sheet.iter_rows(min_row=1, max_row=row_count,
                               min_col=1, max_col=column_count):
        for cell in row:
            prototype_column = (last_template_column if cell.column == column_count
                                else min(cell.column, last_template_column - 2))
            cell._style = copy(template["styles"][
                (1 if cell.row == 1 else 2, prototype_column)])
            if cell.row > 1 and cell.column > 1:
                cell.number_format = "0.0000"
    sheet.column_dimensions.clear()
    for column in range(1, column_count + 1):
        letter = get_column_letter(column)
        dim = copy(template["column_a" if column == 1 else "column_data"])
        dim.index = letter
        dim.min = dim.max = column
        sheet.column_dimensions[letter] = dim
    sheet.row_dimensions.clear()
    for row in range(1, row_count + 1):
        dim = copy(template["row_header" if row == 1 else "row_data"])
        dim.index = row
        sheet.row_dimensions[row] = dim


def anonymize_workbook(book) -> None:
    """Remove personal document properties from the submitted result copy."""
    book.properties.creator = ""
    book.properties.lastModifiedBy = ""


def write_fixed_workbook(
    template_name: str,
    output_name: str,
    fields_path: Path,
    include_temperature: bool,
    end_time_s: float | None = None,
) -> dict:
    output_path = RESULTS_DIR / output_name
    shutil.copy2(TEMPLATE_DIR / template_name, output_path)
    book = load_workbook(output_path)
    fields = np.load(fields_path)
    times = fields["times_s"]
    radii_cm = fields["radius_m"] * 100.0
    start = 1 if abs(times[0]) < 1e-12 else 0
    stop = len(times) if end_time_s is None else int(np.searchsorted(times, end_time_s, side="right"))
    written_rows = stop - start + 1
    sheet_variables = (
        [("温度", "temperature_c"), ("水分浓度", "moisture")]
        if include_temperature
        else [(book.sheetnames[0], "moisture")]
    )
    for sheet_name, key in sheet_variables:
        sheet = book[sheet_name]
        template = reset_sheet(sheet)
        sheet.cell(1, 1, template["corner"])
        for column, radius in enumerate(radii_cm, 2):
            sheet.cell(1, column, round(float(radius), 1))
        values = fields[key]
        for index in range(start, stop):
            row_number = index - start + 2
            sheet.cell(row_number, 1, int(round(float(times[index]))))
            for column, value in enumerate(values[index], 2):
                sheet.cell(row_number, column, round(float(value), 4))
        style_sheet(sheet, written_rows, len(radii_cm) + 1, template)
    anonymize_workbook(book)
    book.save(output_path)
    book.close()
    fields.close()
    return {
        "path": str(output_path),
        "rows_including_header": written_rows,
        "columns": len(radii_cm) + 1,
        "sheets": [name for name, _ in sheet_variables],
    }


def write_moving_workbook() -> dict:
    fields = np.load(RESULTS_DIR / "q4" / "fields.npz")
    numerics = json.loads((RESULTS_DIR / "q4" / "summary.json").read_text(encoding="utf-8"))["numerics"]
    native_xi = 1.0 - (1.0 - np.linspace(0.0, 1.0, numerics["nodes"])) ** numerics["grid_power"]
    if fields["xi"].shape != native_xi.shape or not np.allclose(fields["xi"], native_xi, rtol=0, atol=1e-14):
        fields.close()
        raise ValueError("Q4 output requires native calculation nodes; regenerate q4/fields.npz before exporting.")
    output_path = RESULTS_DIR / "result4.xlsx"
    shutil.copy2(TEMPLATE_DIR / "result4.xlsx", output_path)
    book = load_workbook(output_path)
    sheet = book[book.sheetnames[0]]
    template = reset_sheet(sheet)
    times = fields["times_s"]
    xi = fields["xi"]
    radii = fields["radii_m"]
    moisture = fields["moisture"]
    fixed_radii_cm = np.arange(0.0, 2.0, 0.1)
    headers = ([template["corner"]]
               + [round(float(radius), 1) for radius in fixed_radii_cm]
               + ["药材表面"])
    for column, value in enumerate(headers, 1):
        sheet.cell(1, column, value)
    start = 1 if abs(times[0]) < 1e-12 else 0
    for index in range(start, len(times)):
        current_radius_m = float(radii[index])
        profile = moisture[index]
        row: list[float | int | None] = [int(round(float(times[index])))]
        for radius_cm in fixed_radii_cm:
            position_m = radius_cm / 100.0
            if position_m > current_radius_m + 1e-12:
                row.append(None)
            else:
                value = np.interp(position_m / current_radius_m, xi, profile)
                row.append(round(float(value), 4))
        row.append(round(float(profile[-1]), 4))
        for column, value in enumerate(row, 1):
            sheet.cell(index - start + 2, column, value)
    rows = len(times) - start + 1
    columns = len(fixed_radii_cm) + 2
    style_sheet(sheet, rows, columns, template)
    anonymize_workbook(book)
    book.save(output_path)
    book.close()
    fields.close()
    return {
        "path": str(output_path),
        "rows_including_header": rows,
        "columns": columns,
        "sheets": book.sheetnames,
        "blank_rule": "固定物理位置超过当时药材半径时留空；末列始终为实时表面。",
    }


def verify_workbook(item: dict) -> dict:
    path = Path(item["path"])
    book = load_workbook(path, read_only=True, data_only=False)
    verification = {
        "path": str(path),
        "file_size_bytes": path.stat().st_size,
        "sheets": {},
        "formula_errors": [],
    }
    error_tokens = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
    for sheet in book.worksheets:
        first_row = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        last_row = [
            cell.value
            for cell in next(
                sheet.iter_rows(
                    min_row=sheet.max_row,
                    max_row=sheet.max_row,
                    min_col=1,
                    max_col=sheet.max_column,
                )
            )
        ]
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str) and value in error_tokens:
                    verification["formula_errors"].append(
                        {"sheet": sheet.title, "value": value}
                    )
        verification["sheets"][sheet.title] = {
            "rows": sheet.max_row,
            "columns": sheet.max_column,
            "first_row": first_row,
            "last_row": last_row,
        }
    verification["pass"] = (
        not verification["formula_errors"]
        and all(
            details["rows"] == item["rows_including_header"]
            and details["columns"] == item["columns"]
            for details in verification["sheets"].values()
        )
    )
    question = int(path.stem[-1])
    fields_path = RESULTS_DIR / f"q{question}" / "fields.npz"
    with np.load(fields_path) as archive:
        # Decode each compressed array once, rather than once per Excel row.
        fields = {key: archive[key] for key in archive.files}
    times = fields["times_s"]
    end_time = {1: 1800.0, 2: 10800.0}.get(question, float(times[-1]))
    indexes = np.flatnonzero((times > 0) & (times <= end_time))
    interval = 1 if question <= 2 else 60
    expected_times = np.arange(interval, end_time + 0.5, interval)
    schedule_valid = np.array_equal(times[indexes], expected_times)
    mismatches = 0
    for sheet in book.worksheets:
        key = "temperature_c" if sheet.title == "温度" else "moisture"
        if sheet.max_row != len(indexes) + 1:
            mismatches += 1
        for index, row in zip(indexes, sheet.iter_rows(min_row=2, values_only=True)):
            if row[0] != int(times[index]):
                mismatches += 1
            if question != 4:
                values = np.interp(np.linspace(0, 0.02, 21), fields["radius_m"], fields[key][index])
                expected = [round(float(v), 4) for v in values]
            else:
                radius = float(fields["radii_m"][index])
                expected = [None if r > radius + 1e-12 else
                            round(float(np.interp(r / radius, fields["xi"], fields["moisture"][index])), 4)
                            for r in np.arange(20) * 0.001]
                expected.append(round(float(fields["moisture"][index, -1]), 4))
            mismatches += sum(actual != target for actual, target in zip(row[1:], expected))
    verification["source_value_mismatches"] = mismatches
    verification["time_schedule_valid"] = schedule_valid
    verification["source_fields_sha256"] = hashlib.sha256(fields_path.read_bytes()).hexdigest()
    verification["workbook_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    verification["pass"] = verification["pass"] and schedule_valid and mismatches == 0
    book.close()
    return verification


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    outputs = [
        write_fixed_workbook(
            "result1.xlsx",
            "result1.xlsx",
            RESULTS_DIR / "q1" / "fields.npz",
            include_temperature=True,
            end_time_s=1800.0,
        ),
        write_fixed_workbook(
            "result2.xlsx",
            "result2.xlsx",
            RESULTS_DIR / "q2" / "fields.npz",
            include_temperature=True,
            end_time_s=10800.0,
        ),
        write_fixed_workbook(
            "result3.xlsx",
            "result3.xlsx",
            RESULTS_DIR / "q3" / "fields.npz",
            include_temperature=False,
        ),
        write_moving_workbook(),
    ]
    report = {
        "authoring_backend": "existing openpyxl result exporter",
        "template_policy": "保留附件模板表头、工作表名称和样式；展开省略号，结果显示四位小数。",

        "outputs": outputs,
        "verification": [verify_workbook(item) for item in outputs],
    }
    report_path = RESULTS_DIR / "workbook_verification.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if not all(item["pass"] for item in report["verification"]):
        raise RuntimeError("Workbook verification failed; see workbook_verification.json")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
