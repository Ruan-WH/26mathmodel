from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator


ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent


def read_numeric_xlsx(path: Path) -> np.ndarray:
    ws = load_workbook(path, data_only=True, read_only=True).active
    rows = []
    for row in ws.iter_rows(values_only=True):
        try:
            rows.append([float(v) for v in row])
        except (TypeError, ValueError):
            continue
    return np.asarray(rows, dtype=float)


def write_table(path: Path, x: np.ndarray, y: np.ndarray) -> None:
    np.savetxt(path, np.column_stack([x, y]), fmt="%.12g", delimiter=" ")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    env = read_numeric_xlsx(ROOT / "附件" / "附件1.xlsx")
    radius = read_numeric_xlsx(ROOT / "附件" / "附件2.xlsx")

    end_s = 183000.0
    dense_t = np.arange(0.0, end_s + 60.0, 60.0)
    te_interp = PchipInterpolator(env[:, 0], env[:, 1], extrapolate=False)
    ce_interp = PchipInterpolator(env[:, 0], env[:, 2], extrapolate=False)
    te = np.where(dense_t <= env[-1, 0], te_interp(np.minimum(dense_t, env[-1, 0])), env[-1, 1])
    ce = np.where(dense_t <= env[-1, 0], ce_interp(np.minimum(dense_t, env[-1, 0])), env[-1, 2])

    monotone_radius_m = np.minimum.accumulate(radius[:, 1]) / 100.0
    r_interp = PchipInterpolator(radius[:, 0], monotone_radius_m, extrapolate=False)
    rt = r_interp(dense_t)

    write_table(OUT / "ambient_temperature.txt", dense_t, te)
    write_table(OUT / "ambient_moisture.txt", dense_t, ce)
    write_table(OUT / "radius_history.txt", dense_t, rt)
    print(f"Prepared {len(dense_t)} rows through {end_s/3600:.3f} h")
    print(f"Radius: {rt[0]*100:.3f} cm -> {rt[-1]*100:.3f} cm")


if __name__ == "__main__":
    main()
