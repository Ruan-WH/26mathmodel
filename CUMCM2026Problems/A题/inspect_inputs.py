from pathlib import Path

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parent / "附件"


def preview_workbook(path: Path) -> None:
    print(f"\n=== {path.name} ===")
    book = load_workbook(path, read_only=True, data_only=False)
    print("sheets:", book.sheetnames)
    for sheet in book.worksheets:
        print(f"[{sheet.title}] size={sheet.max_row}x{sheet.max_column}")
        rows = list(sheet.iter_rows(values_only=True))
        for row in rows[:8]:
            print(row)
        if len(rows) > 8:
            print("... tail ...")
            for row in rows[-3:]:
                print(row)


if __name__ == "__main__":
    paths = [BASE / "附件1.xlsx", BASE / "附件2.xlsx"]
    paths.extend(sorted((BASE / "附件3").glob("*.xlsx")))
    for workbook_path in paths:
        preview_workbook(workbook_path)
