"""Build and verify the explicitly requested flat, minimal support ZIP.

Only delivery copies change: input lookup, root paths and plot output copying.
Canonical modeling sources and the four submitted workbooks are read-only.
"""
from __future__ import annotations

import ast
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import uuid
import zipfile

import numpy as np
from openpyxl import load_workbook

AUDIT = Path(__file__).resolve().parent
PROJECT = AUDIT.parent.parent
WORKSPACE = PROJECT.parents[1]
DELIVERY = AUDIT / "delivery"
ZIP_PATH = PROJECT / "support_materials_20260913.zip"
NAMES = ("solve_drying.py", "make_figures.py", "make_additional_figures.py",
         "figure_style.py", "run_q4_counterfactual.py")
INSTRUCTIONS = """平铺支撑材料：主模型求解与主结果绘图。
环境：Python 3.11+，numpy、scipy、matplotlib；绘图需安装 Times New Roman
及 SimSun（宋体）字体。本交付副本无需外部输入Excel，附件1/2原始数值由同目录
input_data.py提供；数据来源和原始文件SHA-256写在该模块。
从解压根目录依次执行：
  python solve_drying.py --skip-convergence
  python run_q4_counterfactual.py
  python make_figures.py
  python make_additional_figures.py
第一个命令按原模型从初态计算四问；第二个命令生成第四问比较图所需的同物性固定半径
对照。计算可能耗时较长。--skip-convergence仅跳过主入口附带的时间步长与模型对照检查，
不改变主模型终算设置。运行后自动生成results/和figures/目录；ZIP本身没有目录。
图中的温湿场来自重新计算生成的NPZ，不能仅靠四个结果Excel替代全部内部场。
根目录result1.xlsx至result4.xlsx是原样保留的正式结果，本套命令不会改写这些文件。
这里只提供主模型和主结果图；COMSOL、独立验证图、论文排版与结果模板导出程序
不属于这个平铺包。canonical求解函数未改，只对输入读取和交付目录作适配。"""


def digest(data):
    return hashlib.sha256(data).hexdigest()


def module_doc(text, extra):
    tree = ast.parse(text)
    old = ast.get_docstring(tree, clean=False) or ""
    doc = '"""' + extra + ("\n\n" + old if old else "") + '"""\n'
    if ast.get_docstring(tree) is not None:
        node = tree.body[0]
        lines = text.splitlines(keepends=True)
        return "".join(lines[:node.lineno-1]) + doc + "".join(lines[node.end_lineno:])
    return doc + "\n" + text


def replace_function(text, name, replacement):
    node = next(n for n in ast.parse(text).body if isinstance(n, ast.FunctionDef) and n.name == name)
    lines = text.splitlines(keepends=True)
    return "".join(lines[:node.lineno-1]) + replacement.rstrip() + "\n" + "".join(lines[node.end_lineno:])


def read_original_inputs():
    records = {}
    for name in ("附件1.xlsx", "附件2.xlsx"):
        path = PROJECT / "附件" / name
        book = load_workbook(path, read_only=True, data_only=True)
        rows = [[float(v) for v in row] for row in book.active.iter_rows(min_row=2, values_only=True)
                if row[0] is not None]
        book.close()
        array = np.asarray(rows, dtype=float)
        records[name] = dict(source="附件/" + name, source_sha256=digest(path.read_bytes()),
                             shape=list(array.shape), numeric_array_sha256=digest(array.tobytes()),
                             values=rows)
    return records


def input_module(records):
    metadata = {name: {k: v for k, v in item.items() if k != "values"}
                for name, item in records.items()}
    values = {name: item["values"] for name, item in records.items()}
    return ('"""平铺交付的原始输入数值；不是模型输出或压缩隐藏的其他代码。\n'
            '附件1各行依次为时间、环境温度、环境水分；附件2为时间、半径。\n'
            '数值由题给Excel的活动表第2行起逐行按原求解器float规则提取，未拟合、抽样或改值。\n'
            '原文件和数值数组的SHA-256、形状见SOURCE_INFO；无需题给Excel文件即可求解。\n'
            '运行环境与命令顺序见solve_drying.py开头说明。\n"""\n'
            'from pathlib import Path\nimport numpy as np\n\n'
            'SOURCE_INFO = ' + repr(metadata) + '\n\n'
            '_TABLES = ' + repr(values) + '\n\n'
            'def read_numeric_sheet(name):\n'
            '    """Return a fresh array, preserving every supplied numeric value."""\n'
            '    key = Path(name).name\n'
            '    if key not in _TABLES:\n'
            '        raise FileNotFoundError(f"Unsupported embedded input: {key}")\n'
            '    return np.asarray(_TABLES[key], dtype=float).copy()\n')


def definitions(text):
    return {node.name: ast.dump(node, include_attributes=False)
            for node in ast.parse(text).body
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef))}


def make_copies(records):
    copies, ast_checks = {}, []
    for name in NAMES:
        original = (PROJECT / "code" / name).read_text(encoding="utf-8-sig")
        text = original.replace("ROOT = Path(__file__).resolve().parents[1]",
                                "ROOT = Path(__file__).resolve().parent")
        allowed = set()
        if name == "solve_drying.py":
            text = text.replace("from openpyxl import load_workbook",
                                "from input_data import read_numeric_sheet")
            text = replace_function(text, "_read_numeric_sheet",
                "def _read_numeric_sheet(path: Path) -> np.ndarray:\n"
                "    return read_numeric_sheet(path.name)\n")
            allowed = {"_read_numeric_sheet"}
        elif name == "make_additional_figures.py":
            text = text.replace("import shutil\n", "")
            text = text.replace("PAPER = ROOT / 'paper/cumcm-1.1.0/figures'\n", "")
            text = text.replace("REPORT = ROOT / 'reports/figure_additions_20260911'", "REPORT = OUT")
            text = text.replace("    shutil.copy2(OUT/f'{name}.pdf',PAPER/f'{name}.pdf')\n", "")
            text = text.replace("def main():\n", "def main():\n    OUT.mkdir(parents=True, exist_ok=True)\n")
            allowed = {"save", "main"}
        if name != "figure_style.py":
            text = module_doc(text, INSTRUCTIONS)
        a, b = definitions(original), definitions(text)
        assert set(a) == set(b), name
        changed = sorted(k for k in a if a[k] != b[k])
        assert set(changed) == allowed, (name, changed, allowed)
        ast_checks.append(dict(file=name, definitions=len(a), identical_definitions=len(a)-len(changed),
                               changed_io_definitions=changed, all_other_definitions_ast_identical=True))
        compile(text, name, "exec")
        copies[name] = text.encode("utf-8")
    copies["input_data.py"] = input_module(records).encode("utf-8")
    return copies, ast_checks


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def smoke(zip_bytes, records):
    location = WORKSPACE / "tmp" / "flat_delivery_20260913" / ("run-" + uuid.uuid4().hex[:10])
    location.mkdir(parents=True)
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as archive:
        for name in archive.namelist():
            assert name == Path(name).name and "/" not in name and "\\" not in name
            (location / name).write_bytes(archive.read(name))
    sys.path.insert(0, str(location))
    original = load_module("_flat_audit_original_solver", PROJECT / "code/solve_drying.py")
    flat = load_module("_flat_audit_delivery_solver", location / "solve_drying.py")
    inputs = load_module("_flat_audit_input_values", location / "input_data.py")
    input_checks = []
    for name, item in records.items():
        actual = inputs.read_numeric_sheet(name)
        expected = original._read_numeric_sheet(PROJECT / "附件" / name)
        assert np.array_equal(actual, expected)
        input_checks.append(dict(file=name, shape=list(actual.shape), all_values_identical=True,
                                 value_count=int(actual.size)))
    comparisons = []
    for prop_name in ("property_q1", "property_q23", "property_q4"):
        expected = original.simulate_fixed(original.load_environment(), getattr(original, prop_name),
                                           end_time_s=10, dt_s=1, output_interval_s=1,
                                           stop_at_threshold=False)
        actual = flat.simulate_fixed(flat.load_environment(), getattr(flat, prop_name),
                                    end_time_s=10, dt_s=1, output_interval_s=1, stop_at_threshold=False)
        for key, value in expected.items():
            if isinstance(value, np.ndarray):
                assert np.array_equal(value, actual[key], equal_nan=True), (prop_name, key)
        comparisons.append(dict(model=prop_name, steps=10, output_arrays_identical=True))
    # Shrinking-domain implicit step at both early and late boundary times.
    for time_s in (1.0, 21600.0, 172800.0):
        results = []
        for model in (original, flat):
            radius, rate = model.load_radius_history().values(time_s)
            ambient_t, ambient_c = model.load_environment().values(time_s)
            t = np.linspace(43.0, 49.0, model.NODES)
            c = np.linspace(1.4, 0.7, model.NODES)
            results.append(model.advance_coupled(t, c, 1.0, radius, rate, ambient_t, ambient_c,
                                                model.property_q4, True))
        assert all(np.array_equal(a, b) for a, b in zip(results[0], results[1]))
        comparisons.append(dict(model="moving_domain_step", time_s=time_s, arrays_identical=True))
    sys.path.pop(0)

    # Exercise the delivered plotting entry points with unchanged saved calculation data.
    # These files are test fixtures only and are never included in the flat ZIP.
    for q in ("q1", "q2", "q3", "q4"):
        target = location / "results" / q
        target.mkdir(parents=True)
        for name in ("fields.npz", "summary.json"):
            shutil.copy2(PROJECT / "results" / q / name, target / name)
    shutil.copy2(PROJECT / "results/q4/fixed_radius_counterfactual.npz",
                 location / "results/q4/fixed_radius_counterfactual.npz")
    env = os.environ.copy()
    env["MPLBACKEND"] = "Agg"
    env["MPLCONFIGDIR"] = str(location / "_matplotlib_cache")
    env["PYTHONUTF8"] = "1"
    runs = []
    for name in ("make_figures.py", "make_additional_figures.py"):
        result = subprocess.run([sys.executable, "-B", name], cwd=location, env=env,
                                text=True, encoding="utf-8", stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        (location / (name + ".stdout.txt")).write_text(result.stdout, encoding="utf-8")
        (location / (name + ".stderr.txt")).write_text(result.stderr, encoding="utf-8")
        if result.returncode:
            raise RuntimeError(f"Plot smoke failed: {name}\n{result.stderr[-4000:]}")
        runs.append(dict(script=name, returncode=0))
    plots = sorted(p.name for p in (location / "figures").glob("*") if p.suffix in {".png", ".pdf"})
    assert len(plots) == 18, plots
    return dict(isolated_directory=location.relative_to(WORKSPACE).as_posix(),
                original_input_checks=input_checks, short_solver_checks=comparisons,
                plot_runs=runs, generated_plot_artifacts=plots,
                plot_data_scope="Existing unchanged saved fields used only as isolated test fixtures; no intermediate data in ZIP",
                full_duration_model_rerun=False)


def build():
    DELIVERY.mkdir(parents=True, exist_ok=True)
    protected_paths = [PROJECT / "code" / name for name in NAMES]
    protected_paths += [PROJECT / "results" / f"result{i}.xlsx" for i in range(1, 5)]
    protected_paths += [PROJECT / "附件" / f"附件{i}.xlsx" for i in (1, 2)]
    before = {p.relative_to(PROJECT).as_posix(): digest(p.read_bytes()) for p in protected_paths}
    records = read_original_inputs()
    copies, ast_checks = make_copies(records)
    for name, data in copies.items():
        (DELIVERY / name).write_bytes(data)
    payloads = dict(copies)
    for i in range(1, 5):
        payloads[f"result{i}.xlsx"] = (PROJECT / "results" / f"result{i}.xlsx").read_bytes()
    ai = PROJECT / "paper/AI_usage/AI工具使用详情.pdf"
    payloads[ai.name] = ai.read_bytes()
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name, data in sorted(payloads.items()):
            assert "/" not in name and "\\" not in name
            archive.writestr(name, data)
    raw = output.getvalue()
    assert len(raw) <= 20_000_000
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        assert len(archive.namelist()) == 11 and not any(i.is_dir() for i in archive.infolist())
        assert archive.testzip() is None
        assert all(archive.read(name) == data for name, data in payloads.items())
    smoke_checks = smoke(raw, records)
    after = {p.relative_to(PROJECT).as_posix(): digest(p.read_bytes()) for p in protected_paths}
    assert before == after

    roles = {"solve_drying.py": "四问主模型求解；平铺路径与内嵌输入适配",
             "make_figures.py": "四问主要结果图绘制",
             "make_additional_figures.py": "温湿时空场、扩散率、干燥前沿和收缩主图绘制",
             "figure_style.py": "主图字体、数学排版和缺字检查",
             "run_q4_counterfactual.py": "第四问同物性固定半径对照，供主要比较图使用",
             "input_data.py": "附件1/2原始数值与来源哈希；无需外部输入文件"}
    entries = []
    for name, data in sorted(payloads.items()):
        if name in copies:
            source = "code/" + name if name != "input_data.py" else None
            delivery_source = (DELIVERY / name).relative_to(PROJECT).as_posix()
            kind = "Python"
        elif name.endswith(".xlsx"):
            source = "results/" + name
            delivery_source = source
            kind = "result_workbook"
        else:
            source = ai.relative_to(PROJECT).as_posix()
            delivery_source = source
            kind = "AI_usage_pdf"
        entries.append(dict(archive_path=name, source=source, delivery_source=delivery_source,
                            description=roles.get(name, "原样保留的正式结果工作簿" if kind == "result_workbook" else "AI工具使用情况说明"),
                            kind=kind, size_bytes=len(data), sha256=digest(data),
                            source_sha256=digest((PROJECT/source).read_bytes()) if source else None))
    manifest = dict(schema_version=1, layout="flat", file_count=11, files=entries,
                    source_inputs=[{k:v for k,v in item.items() if k != "values"} for item in records.values()],
                    code_adaptation_scope="Only delivery IO roots, embedded input lookup and output copying; all modeling functions unchanged",
                    directories_in_archive=False, metadata_manifest_in_archive=False)
    checks = dict(status="PASS", zip_bytes=len(raw), zip_sha256=digest(raw), archive_file_count=11,
                  archive_names=sorted(payloads), no_directory_entries=True, all_root_names=True,
                  four_result_workbooks_byte_identical=True, ai_pdf_byte_identical=True,
                  definitions_ast_checks=ast_checks, protected_files_unchanged=True,
                  protected_sha256_before=before, protected_sha256_after=after, **smoke_checks)
    # Only publish after the flat contents and actual delivery scripts pass verification.
    ZIP_PATH.write_bytes(raw)
    (AUDIT / "support_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    (AUDIT / "checks.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps({k:checks[k] for k in ("status", "zip_bytes", "zip_sha256", "archive_names",
                                          "four_result_workbooks_byte_identical", "full_duration_model_rerun")},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    sys.dont_write_bytecode = True
    sys.stdout.reconfigure(encoding="utf-8")
    build()
