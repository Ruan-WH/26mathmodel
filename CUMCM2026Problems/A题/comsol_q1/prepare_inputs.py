"""Prepare COMSOL boundary data and compare its output with the Q1 baseline."""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator


HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent


def read_environment() -> np.ndarray:
    book = load_workbook(PROJECT / "附件" / "附件1.xlsx", read_only=True, data_only=True)
    rows = []
    for row in book.active.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            rows.append(tuple(float(x) for x in row[:3]))
    return np.asarray(rows)


def prepare() -> None:
    raw = read_environment()
    seconds = np.arange(0.0, 1800.0 + 1.0, 1.0)
    temperature = PchipInterpolator(raw[:, 0], raw[:, 1])(seconds)
    moisture = PchipInterpolator(raw[:, 0], raw[:, 2])(seconds)
    for name, values in (("ambient_temperature.txt", temperature), ("ambient_moisture.txt", moisture)):
        with (HERE / name).open("w", encoding="ascii", newline="") as handle:
            writer = csv.writer(handle, delimiter=" ", lineterminator="\n")
            writer.writerows(zip(seconds, values))
    metadata = {
        "source": str(PROJECT / "附件" / "附件1.xlsx"),
        "interpolation": "PCHIP sampled each second; COMSOL uses linear interpolation",
        "time_range_s": [0, 1800],
        "rows": len(seconds),
        "initial": {"temperature_c": float(temperature[0]), "moisture": float(moisture[0])},
        "final": {"temperature_c": float(temperature[-1]), "moisture": float(moisture[-1])},
    }
    (HERE / "input_metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def compare(csv_path: Path) -> None:
    """Compare COMSOL point export with the existing 1D finite-volume baseline."""
    baseline = json.loads((PROJECT / "results" / "q1" / "summary.json").read_text(encoding="utf-8"))
    # COMSOL Data exports contain comment lines beginning with %. Remaining rows are numeric.
    data = np.loadtxt(csv_path, comments="%", delimiter=",")
    if data.ndim == 1:
        data = data[None, :]
    comsol_times = [0, 100, 300, 600, 900, 1200, 1500, 1800]
    selected_radii = [0, 5, 10, 15, 20]
    time_columns = [comsol_times.index(t) for t in baseline["table_times_s"]]
    comsol_t = np.asarray([[data[r, 2 + 2 * j] for r in selected_radii] for j in time_columns])
    comsol_c = np.asarray([[data[r, 3 + 2 * j] for r in selected_radii] for j in time_columns])
    base_t = np.asarray(baseline["temperature_table_c"])
    base_c = np.asarray(baseline["moisture_table"])
    delta_t = comsol_t - base_t
    delta_c = comsol_c - base_c
    result = {
        "comsol_csv": str(csv_path),
        "shape": list(data.shape),
        "baseline_times_s": baseline["table_times_s"],
        "baseline_radii_cm": baseline["table_radii_cm"],
        "max_abs_temperature_difference_c": float(np.max(np.abs(delta_t))),
        "rmse_temperature_c": float(np.sqrt(np.mean(delta_t**2))),
        "max_abs_moisture_difference": float(np.max(np.abs(delta_c))),
        "rmse_moisture": float(np.sqrt(np.mean(delta_c**2))),
        "comsol_final_temperature_c": comsol_t[-1].tolist(),
        "baseline_final_temperature_c": base_t[-1].tolist(),
        "comsol_final_moisture": comsol_c[-1].tolist(),
        "baseline_final_moisture": base_c[-1].tolist(),
    }
    (HERE / "comparison.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    with (HERE / "q1_comparison.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["time_s", "radius_cm", "comsol_T_C", "baseline_T_C", "delta_T_C",
                         "comsol_C", "baseline_C", "delta_C"])
        for i, time_s in enumerate(baseline["table_times_s"]):
            for j, radius_cm in enumerate(baseline["table_radii_cm"]):
                writer.writerow([time_s, radius_cm, comsol_t[i, j], base_t[i, j], delta_t[i, j],
                                 comsol_c[i, j], base_c[i, j], delta_c[i, j]])


if __name__ == "__main__":
    prepare()
    if len(sys.argv) == 2:
        compare(Path(sys.argv[1]))
